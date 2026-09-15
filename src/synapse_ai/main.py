import base64
import binascii
import hashlib
import hmac
import io
import json
import os
import re
import zipfile
from typing import Any, TypedDict

import httpx
from fastapi import Depends, FastAPI, Header, HTTPException
from pydantic import BaseModel, ConfigDict, Field

from .telemetry import setup_telemetry, span

app = FastAPI(title="Synapse private AI", version="1.0.0")
setup_telemetry(app)

# Standard EICAR antivirus test signature (https://www.eicar.org/) — every
# real antivirus engine, and this heuristic scanner, treats a file containing
# it as malicious. It is inert (not real malware) and is the industry way to
# verify a scanner actually rejects something without needing live malware.
EICAR_SIGNATURE = (
    b"X5O!P%@AP[4\\PZX54(P^)7CC)7}$" b"EICAR-STANDARD-ANTIVIRUS-TEST-FILE!$H+H*"
)
# Magic-byte prefixes for compiled executables/libraries. A "source document"
# uploaded for expertise capture should never legitimately be one of these.
_EXECUTABLE_MAGIC: tuple[bytes, ...] = (
    b"MZ",  # Windows PE/EXE/DLL
    b"\x7fELF",  # Linux ELF
    b"\xfe\xed\xfa\xce",  # Mach-O 32-bit
    b"\xfe\xed\xfa\xcf",  # Mach-O 64-bit
    b"\xca\xfe\xba\xbe",  # Mach-O universal binary / Java class
    b"\xcf\xfa\xed\xfe",  # Mach-O 64-bit (reverse byte order)
)


def scan_binary(content: bytes, filename: str) -> None:
    """Heuristic malware/abuse scan for uploaded file bytes (PRD §21
    "uploaded-file scanner"). This is not a substitute for a real antivirus
    engine (e.g. ClamAV) in production — see README — but it deterministically
    rejects the industry-standard EICAR test file, disguised executables, and
    zip/office-document decompression bombs, with zero external dependencies
    or network calls, so it always runs.
    """
    if EICAR_SIGNATURE in content:
        raise HTTPException(
            422, "Rejected: file matches the EICAR antivirus test signature."
        )
    if any(content.startswith(magic) for magic in _EXECUTABLE_MAGIC):
        raise HTTPException(
            422, f"Rejected: '{filename}' is a compiled executable, not a source document."
        )
    if zipfile.is_zipfile(io.BytesIO(content)):
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                expanded = sum(info.file_size for info in archive.infolist())
                compressed = max(1, sum(info.compress_size for info in archive.infolist()))
                if expanded > 500 * 1024 * 1024 or expanded / compressed > 200:
                    raise HTTPException(
                        422, "Rejected: archive expands suspiciously large (decompression-bomb heuristic)."
                    )
        except zipfile.BadZipFile:
            pass
    cmd = os.getenv("CLAMAV_CMD", "").strip()
    required = os.getenv("CLAMAV_REQUIRED", "").lower() == "true"
    if not cmd:
        if required:
            raise HTTPException(
                503,
                "File scanning is required (CLAMAV_REQUIRED=true) but CLAMAV_CMD is not configured.",
            )
        return
    import subprocess
    import tempfile

    with tempfile.NamedTemporaryFile(suffix="-" + os.path.basename(filename)) as tmp:
        tmp.write(content)
        tmp.flush()
        parts = cmd.split()
        try:
            completed = subprocess.run(
                [*parts, tmp.name],
                capture_output=True,
                timeout=30,
                check=False,
            )
        except (FileNotFoundError, subprocess.TimeoutExpired) as err:
            raise HTTPException(
                503,
                "Configured ClamAV scanner could not run. Refusing to accept the upload as clean.",
            ) from err
        if completed.returncode == 1:
            raise HTTPException(422, "Rejected: ClamAV reported malware.")
        if completed.returncode != 0:
            raise HTTPException(
                503,
                "ClamAV scanner returned an error. Retry after the scanning service is healthy.",
            )


class Metric(TypedDict):
    name: str
    score: float
    threshold: float


def authorize(authorization: str = Header(default="")) -> None:
    token = os.getenv("AI_SERVICE_TOKEN", "")
    if len(token) < 32 or not hmac.compare_digest(authorization, f"Bearer {token}"):
        raise HTTPException(401, "Private service authorization required")


class Request(BaseModel):
    model_config = ConfigDict(extra="allow")
    capsuleId: str = Field(min_length=1, max_length=100)
    idempotencyKey: str = Field(min_length=1, max_length=200)
    allowedScope: dict[str, list[str]]


class Item(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str
    sourceId: str
    contributor: str
    kind: str
    text: str
    status: str


class Query(Request):
    query: str = Field(min_length=1, max_length=4000)
    version: str
    knowledge: list[Item] = Field(max_length=10000)


STOP = {
    "what",
    "does",
    "this",
    "that",
    "with",
    "have",
    "from",
    "your",
    "would",
    "should",
    "about",
    "when",
    "where",
    "which",
    "there",
    "their",
}


def retrieve(request: Query) -> list[Item]:
    allowed = set(request.allowedScope.get("knowledgeIds", []))
    words = {
        word
        for word in re.findall(r"[a-z]{4,}", request.query.lower())
        if word not in STOP
    }
    candidates = []
    for item in request.knowledge:
        if item.status != "APPROVED" or item.id not in allowed:
            continue
        tokens = set(re.findall(r"[a-z]{4,}", item.text.lower()))
        overlap = len(words & tokens)
        if overlap and overlap / max(1, len(words)) >= 0.25:
            candidates.append((overlap, item))
    candidates.sort(key=lambda entry: entry[0], reverse=True)
    return [item for _, item in candidates[:3]]


def answer(request: Query) -> dict[str, Any]:
    with span("query.answer", capsuleId=request.capsuleId, version=request.version):
        return _answer(request)


def _answer(request: Query) -> dict[str, Any]:
    matches = retrieve(request)
    if not matches:
        return {
            "text": "This capsule does not contain approved knowledge that supports an answer to this question.",
            "abstained": True,
            "citations": [],
            "version": request.version,
            "method": "extractive-v1",
        }
    return {
        "text": "The capsule’s approved knowledge indicates:\n\n"
        + "\n\n".join(
            f"{item.text} [{index + 1}]" for index, item in enumerate(matches)
        ),
        "abstained": False,
        "version": request.version,
        "method": "extractive-v1",
        "citations": [
            {
                "id": item.sourceId,
                "source": f"Approved source {item.sourceId[:8]}",
                "quote": item.text,
                "contributor": item.contributor,
                "knowledgeId": item.id,
            }
            for item in matches
        ],
    }


@app.get("/internal/v1/health", dependencies=[Depends(authorize)])
def health() -> dict[str, str]:
    return {"status": "ok", "retrieval": "extractive-v1", "schemaVersion": "1.0"}


@app.post("/internal/v1/query", dependencies=[Depends(authorize)])
def query(request: Query) -> dict[str, Any]:
    return answer(request)


class Extract(Request):
    sourceId: str
    text: str = Field(min_length=20, max_length=1000000)


@app.post("/internal/v1/extract", dependencies=[Depends(authorize)])
def extract(request: Extract) -> dict[str, Any]:
    with span("extract.candidates", capsuleId=request.capsuleId, sourceId=request.sourceId):
        return _extract(request)


def _extract(request: Extract) -> dict[str, Any]:
    if request.sourceId not in request.allowedScope.get("sourceIds", []):
        raise HTTPException(403, "Source is outside authorized scope")
    sentences = [
        s.strip()
        for s in re.split(r"(?<=[.!?])\s+|\n+", request.text)
        if len(s.strip()) > 20
    ]
    items = []
    for index, sentence in enumerate(sentences[:200]):
        lower = sentence.lower()
        kind = (
            "EXCEPTION"
            if any(w in lower for w in ["unless", "except", "only if", "safety risk"])
            else "PROCEDURE"
            if any(w in lower for w in ["first,", "then,", "step "])
            else "HEURISTIC"
            if any(w in lower for w in ["usually", "before", "check"])
            else "CLAIM"
        )
        items.append(
            {
                "kind": kind,
                "text": sentence,
                "quote": sentence,
                "segmentRef": f"{request.sourceId}:{index}",
                "confidence": 1.0,
                "critical": False,
            }
        )
    return {
        "items": items,
        "method": "verbatim-candidate-extraction-v1",
        "confidenceMeaning": "Exact source text match, not truth or expert approval",
    }


class Document(Request):
    sourceId: str
    contentType: str
    filename: str
    contentBase64: str = Field(max_length=140000000)


def ocr_pdf_pages(content: bytes, max_pages: int = 500) -> str:
    """Renders each PDF page to a bitmap (via bundled PDFium — no system
    poppler dependency) and runs local Tesseract OCR over it. Used only when
    a PDF has zero extractable text layer (i.e. it is a scan/photo), per PRD
    §21's "image OCR" gap. Requires the `tesseract` binary on PATH; see
    README/Dockerfile for how it is installed. Never fabricates text: if
    Tesseract is not installed, this raises a clear 503 rather than
    returning empty or invented content.
    """
    import pypdfium2 as pdfium

    try:
        import pytesseract
    except ImportError as err:  # pragma: no cover - always installed per pyproject
        raise HTTPException(503, "OCR library is not installed.") from err
    try:
        pytesseract.get_tesseract_version()
    except OSError as err:
        raise HTTPException(
            503,
            "OCR requires the 'tesseract' binary, which is not installed on "
            "this host. Install tesseract-ocr (see README) or configure a "
            "document-extraction provider.",
        ) from err
    pdf = pdfium.PdfDocument(content)
    try:
        if len(pdf) > max_pages:
            raise HTTPException(413, "Document exceeds page limit")
        pages_text: list[str] = []
        for page in pdf:
            bitmap = page.render(scale=2.0)
            image = bitmap.to_pil()
            pages_text.append(pytesseract.image_to_string(image))
        return "\n".join(pages_text)
    finally:
        pdf.close()


async def _extract_document_text(request: Document) -> str:
    content = base64.b64decode(request.contentBase64, validate=True)
    if len(content) > 100 * 1024 * 1024:
        raise HTTPException(413, "Source too large")
    scan_binary(content, request.filename)
    if request.contentType.startswith("text/"):
        text = content.decode("utf-8")
    elif request.filename.lower().endswith(".pdf"):
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(content))
        if len(reader.pages) > 500:
            raise HTTPException(413, "Document exceeds page limit")
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
        if not text.strip():
            with span("document.ocr_fallback", sourceId=request.sourceId):
                text = ocr_pdf_pages(content)
    elif request.filename.lower().endswith(".docx"):
        from docx import Document as WordDocument

        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 100 * 1024 * 1024:
                raise HTTPException(413, "Expanded document too large")
        doc = WordDocument(io.BytesIO(content))
        text = "\n".join(p.text for p in doc.paragraphs)
    elif request.contentType.startswith(("audio/", "video/")):
        url = os.getenv("TRANSCRIPTION_API_URL")
        key = os.getenv("TRANSCRIPTION_API_KEY")
        if not url or not key:
            raise HTTPException(503, "Transcription provider is not configured")
        async with httpx.AsyncClient(timeout=120) as client:
            result = await client.post(
                url,
                headers={"Authorization": f"Bearer {key}"},
                files={"file": (request.filename, content, request.contentType)},
                data={"model": os.getenv("TRANSCRIPTION_MODEL", "")},
            )
            result.raise_for_status()
            text = result.json()["text"]
    else:
        raise HTTPException(
            415, "This source format needs a configured extraction provider"
        )
    if not text.strip():
        raise HTTPException(
            422, "No readable text found; scanned documents require OCR"
        )
    return text[:1000000]


@app.post("/internal/v1/document", dependencies=[Depends(authorize)])
async def document(request: Document) -> dict[str, str]:
    if request.sourceId not in request.allowedScope.get("sourceIds", []):
        raise HTTPException(403, "Source outside authorized scope")
    with span(
        "document.extract", sourceId=request.sourceId, contentType=request.contentType
    ):
        try:
            return {"text": await _extract_document_text(request)}
        except HTTPException:
            raise
        except (
            ValueError,
            KeyError,
            TypeError,
            OSError,
            binascii.Error,
            zipfile.BadZipFile,
            httpx.HTTPError,
        ):
            raise HTTPException(422, "The source could not be safely parsed") from None


class Interview(Request):
    transcript: str = Field(default="", max_length=100000)
    coverage: list[str] = Field(default_factory=list)


@app.post("/internal/v1/interview/next", dependencies=[Depends(authorize)])
def next_question(request: Interview) -> dict[str, Any]:
    topics = [
        ("decisions", "Describe a difficult decision. What evidence did you trust?"),
        ("exceptions", "When would your usual approach be the wrong one?"),
        (
            "failures",
            "Tell me about a failure. What warning sign did you recognise afterward?",
        ),
        ("judgement", "What do you notice that someone new to the work might miss?"),
        ("boundaries", "When should someone stop and ask for specialist help?"),
    ]
    topic, question = next(
        (
            (topic, question)
            for topic, question in topics
            if topic not in request.coverage
        ),
        (
            "counterexamples",
            "What example would challenge the guidance you have just given?",
        ),
    )
    return {
        "question": question,
        "coverage": [*request.coverage, topic],
        "method": "coverage-planner-v1",
    }


class GoldenCase(BaseModel):
    question: str = Field(min_length=5, max_length=4000)
    expectedElements: list[str] = Field(default_factory=list)
    forbiddenElements: list[str] = Field(default_factory=list)
    unsupported: bool = False


class Evaluation(Request):
    knowledge: list[Item]
    goldenCases: list[GoldenCase] = Field(default_factory=list, max_length=500)


@app.post("/internal/v1/evals/run", dependencies=[Depends(authorize)])
def evaluate(request: Evaluation) -> dict[str, Any]:
    approved = [
        item
        for item in request.knowledge
        if item.status == "APPROVED"
        and item.id in request.allowedScope.get("knowledgeIds", [])
    ]
    cases = []
    coverage_hits = 0
    coverage_total = 0
    unsupported_hits = 0
    unsupported_total = 0
    citation_total = 0
    citation_valid = 0
    for case in request.goldenCases:
        result = answer(
            Query(
                capsuleId=request.capsuleId,
                idempotencyKey=request.idempotencyKey,
                allowedScope=request.allowedScope,
                query=case.question,
                version="candidate",
                knowledge=approved,
            )
        )
        content = result["text"].lower()
        hits = sum(element.lower() in content for element in case.expectedElements)
        if case.unsupported:
            unsupported_total += 1
            unsupported_hits += int(result["abstained"])
        else:
            coverage_total += len(case.expectedElements)
            coverage_hits += hits
        for citation in result["citations"]:
            citation_total += 1
            citation_valid += int(
                any(
                    item.id == citation["knowledgeId"]
                    and item.sourceId == citation["id"]
                    and item.text == citation["quote"]
                    and item.contributor == citation["contributor"]
                    for item in approved
                )
            )
        correct = (
            result["abstained"]
            if case.unsupported
            else (
                not result["abstained"]
                and bool(case.expectedElements)
                and hits == len(case.expectedElements)
            )
        )
        correct = correct and not any(
            element.lower() in content for element in case.forbiddenElements
        )
        cases.append(
            {
                "question": case.question,
                "expected": "Abstain"
                if case.unsupported
                else "; ".join(case.expectedElements),
                "passed": correct,
            }
        )
    valid = 100 * citation_valid / citation_total if citation_total else 0
    metrics: list[Metric] = [
        {"name": "Citation target validity", "score": valid, "threshold": 95.0},
        {
            "name": "Required-element coverage",
            "score": 100 * coverage_hits / coverage_total if coverage_total else 0,
            "threshold": 90.0,
        },
        {
            "name": "Unsupported-question abstention",
            "score": 100 * unsupported_hits / unsupported_total
            if unsupported_total
            else 0,
            "threshold": 100.0,
        },
        {"name": "No fabricated citations", "score": valid, "threshold": 100.0},
    ]
    suite_hash = hashlib.sha256(
        json.dumps(
            [case.model_dump() for case in request.goldenCases], sort_keys=True
        ).encode()
    ).hexdigest()
    return {
        "passed": bool(
            approved
            and coverage_total
            and unsupported_total
            and all(metric["score"] >= metric["threshold"] for metric in metrics)
            and all(case["passed"] for case in cases)
        ),
        "metrics": metrics,
        "cases": cases,
        "suiteHash": suite_hash,
        "method": "expert-golden-extractive-v1",
    }


class GenerateEvals(Request):
    knowledge: list[Item] = Field(default_factory=list, max_length=1000)


@app.post("/internal/v1/evals/generate", dependencies=[Depends(authorize)])
def generate_evals(request: GenerateEvals) -> dict[str, Any]:
    approved = [
        item
        for item in request.knowledge
        if item.status == "APPROVED"
        and item.id in request.allowedScope.get("knowledgeIds", [])
    ]
    generated_cases: list[dict[str, Any]] = []

    # Generate positive supported golden cases from approved knowledge
    for item in approved[:15]:
        text = item.text.strip()
        # Extract meaningful expected key phrases (at least 3 words)
        words = text.split()
        if len(words) >= 3:
            key_phrase = " ".join(words[: min(6, len(words))])
        else:
            key_phrase = text

        if item.kind == "PROCEDURE":
            q = f"How should this procedure be carried out according to the capsule: {key_phrase}?"
        elif item.kind == "HEURISTIC":
            q = f"What guidance does the capsule provide regarding {key_phrase}?"
        elif item.kind == "EXCEPTION":
            q = f"Under what conditions does the following exception apply: {key_phrase}?"
        else:
            q = f"What does the approved knowledge indicate regarding {key_phrase}?"

        generated_cases.append(
            {
                "question": q,
                "expectedElements": [key_phrase],
                "forbiddenElements": ["unsupported invented assumption"],
                "unsupported": False,
            }
        )

    # Generate negative out-of-scope golden cases that must trigger abstention
    unsupported_questions = [
        "What is the best recipe for baking chocolate brownies?",
        "Who was the prime minister of the United Kingdom in 1923?",
        "How do you repair a mechanical clock from the eighteenth century?",
    ]
    for uq in unsupported_questions:
        generated_cases.append(
            {
                "question": uq,
                "expectedElements": [],
                "forbiddenElements": [],
                "unsupported": True,
            }
        )

    return {
        "cases": generated_cases,
        "count": len(generated_cases),
        "method": "synthesized-golden-evals-v1",
    }


class TranscribeRequest(Request):
    contentBase64: str = Field(max_length=140000000)
    filename: str = Field(default="audio.mp3", max_length=200)
    contentType: str = Field(default="audio/mp3", max_length=100)


@app.post("/internal/v1/transcribe", dependencies=[Depends(authorize)])
async def transcribe(request: TranscribeRequest) -> dict[str, Any]:
    url = os.getenv("TRANSCRIPTION_API_URL")
    key = os.getenv("TRANSCRIPTION_API_KEY")
    content = base64.b64decode(request.contentBase64, validate=True)
    scan_binary(content, request.filename)
    if url and key:
        async with httpx.AsyncClient(timeout=120) as client:
            result = await client.post(
                url,
                headers={"Authorization": f"Bearer {key}"},
                files={"file": (request.filename, content, request.contentType)},
                data={"model": os.getenv("TRANSCRIPTION_MODEL", "")},
            )
            result.raise_for_status()
            text = result.json().get("text", "")
            return {"text": text, "segments": [{"text": text, "start": 0, "end": 10}]}
    # Local-development / test adapter only: a caller (tests, demo fixtures) may
    # upload literal UTF-8 text as a stand-in for audio so the pipeline can be
    # exercised without a real provider. This is never used for genuine binary
    # audio, which always fails UTF-8 decoding and falls through to the error
    # below rather than fabricating a plausible-looking transcript.
    if os.getenv("SYNAPSE_ENV", "production") != "production":
        try:
            decoded = content.decode("utf-8")
            return {
                "text": decoded,
                "segments": [{"text": decoded, "start": 0, "end": 10}],
            }
        except UnicodeDecodeError:
            pass
    raise HTTPException(
        503,
        "Transcription provider is not configured. Set TRANSCRIPTION_API_URL and "
        "TRANSCRIPTION_API_KEY. Refusing to fabricate a transcript.",
    )

